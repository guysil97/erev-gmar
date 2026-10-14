"""Stage 6: build the slideshow from order.csv.

16:9, black background, one picture per slide, fitted without crop/distortion and centered, no text,
auto-advance every 3 s via <p:transition advTm="3000"/> injected into each slide's XML.
Usage: python3 stage6_pptx.py [--order build/order.csv] [--out build/graduation_slideshow.pptx]
"""
import argparse, csv
from pathlib import Path
from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from pptx.util import Emu
from common import OUT as BUILD

ADV_MS = 3000
ap = argparse.ArgumentParser()
ap.add_argument("--order", default=str(BUILD / "order.csv"))
ap.add_argument("--out", default=str(BUILD / "graduation_slideshow.pptx"))
args = ap.parse_args()
base = Path(args.order).parent

prs = Presentation()
prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)   # 13.333" x 7.5" = 16:9
SW, SH = prs.slide_width, prs.slide_height
blank = prs.slide_layouts[6]


def add_transition(slide):
    """Insert <p:transition advTm=.../> right after <p:clrMapOvr> (schema order: cSld, clrMapOvr, transition, timing)."""
    sld = slide._element
    for old in sld.findall(qn("p:transition")):
        sld.remove(old)
    tr = etree.Element(qn("p:transition"))
    tr.set("advTm", str(ADV_MS))
    anchor = sld.find(qn("p:clrMapOvr"))
    if anchor is None:
        anchor = sld.find(qn("p:cSld"))
    anchor.addnext(tr)


rows = list(csv.DictReader(open(args.order, encoding="utf-8-sig")))
for r in rows:
    img = base / r["קובץ"]
    w, h = Image.open(img).size
    scale = min(SW / w, SH / h)
    pw, ph = int(w * scale), int(h * scale)
    slide = prs.slides.add_slide(blank)
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(0, 0, 0)
    slide.shapes.add_picture(str(img), (SW - pw) // 2, (SH - ph) // 2, pw, ph)
    add_transition(slide)

prs.save(args.out)
print(f"slides: {len(rows)} -> {args.out}")
