"""Stage 7: reopen the pptx and verify it; write report.md.

Checks: slide count == order.csv rows, each slide has exactly one picture (in order), no text,
16:9, each slide has <p:transition advTm="3000"> placed right after clrMapOvr.
Usage: python3 stage7_verify.py [--order ...] [--pptx ...] [--manifest ...] [--inventory ...] [--check ...] [--report ...]
"""
import argparse, csv, hashlib
from collections import Counter
from pathlib import Path
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn
from common import OUT as BUILD, INVENTORY, MANIFEST

ap = argparse.ArgumentParser()
ap.add_argument("--order", default=str(BUILD / "order.csv"))
ap.add_argument("--pptx", default=str(BUILD / "graduation_slideshow.pptx"))
ap.add_argument("--manifest", default=str(MANIFEST))
ap.add_argument("--inventory", default=str(INVENTORY))
ap.add_argument("--check", default=str(BUILD / "order_check.md"))
ap.add_argument("--report", default=str(BUILD / "report.md"))
ap.add_argument("--title", default="דוח מצגת ערב הגמר")
args = ap.parse_args()

order = list(csv.DictReader(open(args.order, encoding="utf-8-sig")))
base = Path(args.order).parent
prs = Presentation(args.pptx)
errors = []
if abs(prs.slide_width / prs.slide_height - 16 / 9) > 0.001:
    errors.append("היחס אינו 16:9")
slides = list(prs.slides)
if len(slides) != len(order):
    errors.append(f"מספר שקופיות {len(slides)} != {len(order)} שורות ב-order.csv")
for n, (s, r) in enumerate(zip(slides, order), 1):
    pics = [sh for sh in s.shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE]
    if len(pics) != 1:
        errors.append(f"שקופית {n}: {len(pics)} תמונות")
    elif hashlib.sha1(pics[0].image.blob).digest() != hashlib.sha1((base / r["קובץ"]).read_bytes()).digest():
        errors.append(f"שקופית {n}: התמונה לא תואמת ל-{r['קובץ']}")
    if any(sh.has_text_frame and sh.text_frame.text.strip() for sh in s.shapes):
        errors.append(f"שקופית {n}: יש טקסט")
    tr = s._element.find(qn("p:transition"))
    if tr is None or tr.get("advTm") != "3000":
        errors.append(f"שקופית {n}: חסר מעבר של 3 שניות")
    elif tr.getprevious() is None or tr.getprevious().tag != qn("p:clrMapOvr"):
        errors.append(f"שקופית {n}: המעבר לא במיקום הנכון ב-XML")
print("verification:", "OK" if not errors else f"{len(errors)} errors")
for e in errors[:20]:
    print("  ", e)

# ---------- report ----------
seed = order[0]["seed"] if order else ""
L = [f"# {args.title}", "", f"- קובץ: `{Path(args.pptx).name}`", f"- שקופיות: {len(slides)}  "
     f"(משך: {len(slides) * 3 // 60}:{len(slides) * 3 % 60:02d} דקות, 3 שניות לשקופית)", f"- seed: {seed}",
     f"- בדיקה סופית: {'עברה ✅' if not errors else 'נכשלה ❌'}"]
L += [f"  - {e}" for e in errors]
for title, col in (("לפי מסלול", "מסלול"), ("לפי שולח", "שולח"), ("לפי קטגוריה", "קטגוריה")):
    c = Counter(r[col] or "—" for r in order)
    L += ["", f"## {title}", "", "| | שקופיות |", "|---|---|", *[f"| {k} | {v} |" for k, v in c.most_common()]]
L += ["", "## קבצים שהוחרגו"]
inv = Path(args.inventory)
if inv.exists():
    ex = [r for r in csv.DictReader(open(inv, encoding="utf-8-sig")) if r["status"] != "kept" and not r["reason"].startswith("zip")]
    L += ["", f"סה\"כ {len(ex)}.", "", "| קובץ | סיבה |", "|---|---|"]
    L += [f"| {r['logical']} | {r['reason']}{' של ' + r['dup_of'] if r['dup_of'] else ''} |" for r in ex]
in_order = {r["קובץ"] for r in order}
not_used = [r for r in csv.DictReader(open(args.manifest, encoding="utf-8-sig")) if r["קובץ"] not in in_order]
if not_used:
    L += ["", f"תמונות מה-manifest שלא נכנסו לסדר: {len(not_used)}"]
L += ["", "## אילוצים שהופרו", ""]
chk = Path(args.check)
L += chk.read_text(encoding="utf-8").split("\n")[2:] if chk.exists() else ["(אין דוח בדיקת סדר)"]
Path(args.report).write_text("\n".join(L) + "\n", encoding="utf-8")
print("report ->", args.report)
